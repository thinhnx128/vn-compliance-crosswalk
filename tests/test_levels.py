import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from vncc import levels
from vncc.cli import app
from vncc.export import export_xlsx
from vncc.repository import Repository

runner = CliRunner()

CASES = [
    # (mô tả, câu trả lời, cấp độ mong đợi, tiêu chí phải có)
    ("nội bộ, công cộng", {"system_type": "noi_bo", "info": ["cong_cong"]}, 1, "D11"),
    (
        "nội bộ, có thông tin cá nhân",
        {"system_type": "noi_bo", "info": ["cong_cong", "ca_nhan"]},
        2,
        "D12.K1",
    ),
    (
        "hạ tầng một tổ chức",
        {"system_type": "ha_tang", "info": ["rieng"], "infra_scope": "mot_to_chuc"},
        2,
        "D12.K3",
    ),
    (
        "dịch vụ trực tuyến thường, ít chủ thể",
        {"system_type": "phuc_vu", "info": ["ca_nhan"], "online": True, "pd_basic": 5000},
        2,
        "D12.K2.b",
    ),
    (
        "dịch vụ trực tuyến, 100.000 chủ thể cơ bản",
        {"system_type": "phuc_vu", "info": ["ca_nhan"], "online": True, "pd_basic": 100000},
        3,
        "D13.K2.c-1",
    ),
    (
        "dịch vụ trực tuyến, 10.000 chủ thể nhạy cảm",
        {"system_type": "phuc_vu", "info": ["ca_nhan"], "online": True, "pd_sensitive": 10000},
        3,
        "D13.K2.c-2",
    ),
    (
        "ngành nghề kinh doanh có điều kiện",
        {"system_type": "phuc_vu", "info": ["rieng"], "online": True, "conditional_business": True},
        3,
        "D13.K2.a",
    ),
    ("thủ tục hành chính", {"system_type": "phuc_vu", "info": ["ca_nhan"], "tthc": True}, 3, "D13.K2.b"),
    ("có bí mật nhà nước", {"system_type": "noi_bo", "info": ["bmnn"]}, 3, "D13.K1-1"),
    (
        "hạ tầng một tỉnh",
        {"system_type": "ha_tang", "info": ["rieng"], "infra_scope": "bo_nganh_tinh"},
        3,
        "D13.K3",
    ),
    (
        "ICS công trình cấp II",
        {"system_type": "dieu_khien", "info": ["rieng"], "ics_grade": "II"},
        3,
        "D13.K4",
    ),
    ("ICS công trình cấp I", {"system_type": "dieu_khien", "info": ["rieng"], "ics_grade": "I"}, 4, "D14.K3"),
    (
        "QP-AN nghiêm trọng",
        {"system_type": "khac", "info": ["bmnn"], "qpan_impact": "nghiem_trong"},
        4,
        "D14.K1",
    ),
    (
        "hạ tầng toàn quốc 24/7",
        {"system_type": "ha_tang", "info": ["rieng"], "infra_scope": "toan_quoc_24x7"},
        4,
        "D14.K2-2",
    ),
    ("CPĐT quốc gia", {"system_type": "phuc_vu", "info": ["ca_nhan"], "egov_national": True}, 4, "D14.K2-1"),
    (
        "ICS cấp đặc biệt",
        {"system_type": "dieu_khien", "info": ["rieng"], "ics_grade": "dac_biet"},
        5,
        "D15.K4",
    ),
    (
        "dữ liệu quốc gia",
        {"system_type": "noi_bo", "info": ["rieng"], "national_critical_data": True},
        5,
        "D15.K2",
    ),
    ("hạ tầng quốc tế", {"system_type": "ha_tang", "info": ["rieng"], "infra_scope": "quoc_te"}, 5, "D15.K3"),
    ("QP-AN đặc biệt", {"system_type": "khac", "info": ["bmnn"], "qpan_impact": "dac_biet"}, 5, "D15.K1"),
]


@pytest.mark.parametrize(("name", "answers", "expected", "criterion"), CASES, ids=[c[0] for c in CASES])
def test_levels(name, answers, expected, criterion):
    r = levels.determine(answers)
    assert r.level == expected, name
    assert criterion in {m.id for m in r.matches}


def test_every_criterion_is_exercised():
    covered = {c[3] for c in CASES}
    all_ids = {c["id"] for c in levels.criteria_data()["criteria"]}
    assert all_ids - covered <= {"D12.K2.a", "D13.K1-2"}  # được phủ gián tiếp trong các ca khác


def test_highest_level_wins_and_hidden_answers_ignored():
    # Hệ thống nội bộ có thông tin cá nhân; câu trả lời của câu hỏi ẩn (pd_basic) phải bị bỏ qua.
    r = levels.determine({"system_type": "noi_bo", "info": ["ca_nhan"], "pd_basic": 1_000_000})
    assert r.level == 2
    r = levels.determine({"system_type": "noi_bo", "info": ["ca_nhan", "bmnn"]})
    assert r.level == 3 and r.matches[0].level == 3  # sắp xếp cấp độ cao nhất trước


def test_no_match_warns():
    r = levels.determine({"system_type": "khac", "info": []})
    assert r.level is None
    assert any("Khung quản lý rủi ro" in w for w in r.warnings)


@pytest.mark.parametrize(
    "bad",
    [{"system_type": "x"}, {"info": ["cong_cong"]}, {"system_type": "noi_bo", "pd_basic": -1}, {"foo": 1}],
)
def test_invalid_answers(bad):
    with pytest.raises(ValueError):
        levels.determine(bad)


def test_cli_level_file(tmp_path):
    f = tmp_path / "a.yaml"
    f.write_text("system_type: ha_tang\ninfo: [rieng]\ninfra_scope: bo_nganh_tinh\n", encoding="utf-8")
    res = runner.invoke(app, ["level", "--answers", str(f), "--json"])
    assert res.exit_code == 0, res.output
    assert json.loads(res.output)["level"] == 3


def test_cli_level_interactive():
    # loại hệ thống = 1 (nội bộ); thông tin = 1,3 (công cộng, cá nhân); QP-AN = 1 (không); dữ liệu quốc gia = n
    res = runner.invoke(app, ["level"], input="1\n1,3\n1\nn\n")
    assert res.exit_code == 0, res.output
    assert "Cấp độ sơ bộ: 2" in res.output


def test_site_generation(tmp_path):
    out = tmp_path / "index.html"
    assert runner.invoke(app, ["site", "-o", str(out)]).exit_code == 0
    html = out.read_text(encoding="utf-8")
    assert "__CRITERIA__" not in html and '"D15.K4"' in html


def test_export_xlsx(tmp_path):
    repo = Repository.load()
    out = export_xlsx(repo, "iso27001-2022__nd331-2026", tmp_path / "pack.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames[:2] == ["Hướng dẫn", "Tự đánh giá"]
    assert "Tiêu chí cấp độ" in wb.sheetnames and "Độ phủ" in wb.sheetnames
    sa = wb["Tự đánh giá"]
    assert sa["C2"].value == 2
    assert str(sa["F17"].value).startswith("=IF(INDEX(")
    n = sum(1 for _ in repo.catalog("nd331-2026").controls())
    assert sa.max_row == 16 + n
    assert wb["Đối chiếu"].max_row == 1 + len(repo.mappings["iso27001-2022__nd331-2026"].entries)


SITE = Path(__file__).resolve().parents[1] / "docs" / "cap-do" / "index.html"


def test_published_site_matches_criteria():
    """Trang web trong docs/ phải được sinh lại mỗi khi bộ tiêu chí thay đổi."""
    if not SITE.exists():
        pytest.skip("chưa sinh docs/cap-do/index.html")
    assert levels.criteria_json().replace("</", "<\\/") in SITE.read_text(encoding="utf-8")
