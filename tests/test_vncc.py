from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook
from pydantic import ValidationError
from typer.testing import CliRunner

from vncc.cli import app
from vncc.metrics import tally
from vncc.models import Catalog, ControlState, MappingEntry
from vncc.reports import assessment_template, gap_markdown, write_soa
from vncc.repository import Repository, load_file

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "examples" / "abc-corp"
runner = CliRunner()


@pytest.fixture(scope="module")
def repo():
    return Repository.load(EXAMPLE)


@pytest.fixture(scope="module")
def assessment():
    return load_file(EXAMPLE / "assessment.yaml")


def test_iso_catalog_has_93_controls(repo):
    iso = repo.catalog("iso27001-2022")
    sizes = {g.id: len(g.controls) for g in iso.groups}
    assert sizes == {"5": 37, "6": 8, "7": 14, "8": 34}
    assert sum(sizes.values()) == 93


def test_all_bundled_and_example_data_valid(repo):
    assert repo.validate() == []


def test_catalog_status(repo):
    assert repo.catalog("luat91-2025").status == "skeleton"
    nd = repo.catalog("nd331-2026")
    assert nd.status == "draft" and nd.document_number == "331/2026/NĐ-CP"
    assert str(nd.effective_date) == "2026-08-19"


def test_nd331_unverified_items_are_tagged(repo):
    nd = repo.catalog("nd331-2026").index()
    for cid in ("D33", "D35", "D36", "D37"):
        assert "can-doi-chieu" in nd[cid].tags


def test_level_filter(repo):
    nd = repo.catalog("nd331-2026")
    lv2, lv4, lv5 = (nd.for_level(n).index() for n in (2, 4, 5))
    assert "D30.K5.a" in lv2 and "D30.K5.a" not in lv4
    assert "D30.K8" not in lv2 and "D30.K8" in lv4 and "D30.K8" not in lv5
    assert "D22.K5" in lv5 and "D22.K5" not in lv2
    assert "D31.K2.d" in lv2 and "D31.K2.d" in lv5  # không gắn thẻ cấp độ thì áp dụng mọi cấp độ
    assert nd.for_level(None) is nd


def test_coverage(repo):
    cov = repo.coverage("iso27001-2022__nd331-2026")
    uncovered = {c.id for _, c in cov.target_uncovered}
    assert "D21" in uncovered  # hồ sơ cấp độ: nghĩa vụ hành chính, không có biện pháp ISO tương ứng
    assert "D31.K2.d" not in uncovered
    assert cov.target_covered + len(cov.target_uncovered) == cov.target_total


def test_nd331_level2_assessment(repo):
    a = load_file(EXAMPLE / "assessment-nd331-cap-do-2.yaml")
    assert a.level == 2
    errors, warnings = repo.validate_assessment(a)
    assert errors == []
    assert "D30.K8" not in warnings[0]  # yêu cầu cấp độ 3–4 không bị đòi trong hồ sơ cấp độ 2
    md = gap_markdown(repo, a)
    assert "cấp độ 2" in md
    assert "| D31.K2.d |" in md and "D30.K9" not in md


def test_duplicate_control_rejected():
    with pytest.raises(ValidationError):
        Catalog.model_validate(
            {
                "id": "x",
                "title": "x",
                "version": "1",
                "publisher": "x",
                "status": "draft",
                "groups": [
                    {
                        "id": "g",
                        "title": "g",
                        "controls": [{"id": "1", "title": "a"}, {"id": "1", "title": "b"}],
                    }
                ],
            }
        )


def test_mapping_entry_rules():
    with pytest.raises(ValidationError):
        MappingEntry(
            source="5.1", relationship="intersects", rationale="thiếu target ở đây", confidence="low"
        )
    with pytest.raises(ValidationError):
        MappingEntry(
            source="5.1",
            target="X",
            relationship="no_relationship",
            rationale="không được có target",
            confidence="low",
        )


def test_not_applicable_requires_justification():
    with pytest.raises(ValidationError):
        ControlState(applicable=False)


def test_related_both_directions(repo):
    fwd = repo.related("iso27001-2022", "5.17")
    assert [(r.catalog, r.control, r.relationship) for r in fwd] == [("abc-policy", "CS-02", "subset_of")]
    back = repo.related("abc-policy", "CS-02")
    assert {(r.control, r.relationship) for r in back} == {("5.17", "superset_of"), ("5.18", "intersects")}


def test_invalid_reference_detected(repo, tmp_path):
    bad = tmp_path / "mappings"
    bad.mkdir()
    (bad / "bad.yaml").write_text(
        yaml.safe_dump(
            {
                "kind": "mapping",
                "id": "bad",
                "source_catalog": "iso27001-2022",
                "target_catalog": "abc-policy",
                "version": "0",
                "status": "draft",
                "entries": [
                    {
                        "source": "9.99",
                        "target": "CS-01",
                        "relationship": "equivalent",
                        "rationale": "mã nguồn không tồn tại",
                        "confidence": "low",
                    }
                ],
            },
            allow_unicode=True,
        ),
        encoding="utf-8",
    )
    r = Repository.load(EXAMPLE, tmp_path)
    assert any("9.99" in e for e in r.validate())


def test_tally_rule(repo, assessment):
    t = tally(repo.catalog("iso27001-2022"), assessment)["*"]
    assert t.total == 93 and t.not_applicable == 2 and t.applicable == 91
    c = t.counts
    assert (c["implemented"], c["partial"], c["planned"], c["not_implemented"]) == (6, 4, 2, 1)
    assert c["unassessed"] == 93 - 15
    assert t.score == pytest.approx((6 + 0.5 * 4) / 91)


def test_soa_workbook(repo, assessment, tmp_path):
    out = write_soa(repo, assessment, tmp_path / "soa.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Thông tin", "SoA", "Tổng hợp"]
    ws = wb["SoA"]
    assert ws.max_row == 94
    row = {c.value: i for i, c in enumerate(ws[1])}
    r84 = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "8.4")
    assert r84[row["Áp dụng"]] == "Không"
    r517 = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "5.17")
    assert "abc-policy:CS-02" in r517[row["Yêu cầu liên quan"]]
    assert str(wb["Tổng hợp"]["J6"].value).startswith("=IF(")


def test_gap_report(repo, assessment):
    md = gap_markdown(repo, assessment)
    assert "Ngày chốt số liệu:** 2026-10-31" in md
    assert "| 5.18 | Quyền truy cập | Triển khai một phần" in md
    assert "8.4 Truy cập mã nguồn:** Tổ chức không tự phát triển" in md
    assert "| 5.1 |" not in md  # đã triển khai thì không nằm trong danh sách khoảng cách
    assert "8,8%" in md  # (6 + 0,5 x 4) / 91


def test_template_roundtrip(repo, tmp_path):
    path = tmp_path / "a.yaml"
    path.write_text(assessment_template(repo.catalog("iso27001-2022"), "ABC", "Toàn bộ"), "utf-8")
    a = load_file(path)
    assert len(a.controls) == 93
    assert repo.validate_assessment(a) == ([], [])


def test_cli_end_to_end(tmp_path):
    d = ["-d", str(EXAMPLE)]
    assert runner.invoke(app, ["validate", str(EXAMPLE / "assessment.yaml"), *d]).exit_code == 0
    res = runner.invoke(app, ["list", *d])
    assert "iso27001-2022" in res.output and "abc-policy" in res.output
    res = runner.invoke(app, ["show", "iso27001-2022", "5.24", *d])
    assert "CS-04" in res.output
    out = tmp_path / "soa.xlsx"
    assert runner.invoke(app, ["soa", str(EXAMPLE / "assessment.yaml"), "-o", str(out), *d]).exit_code == 0
    assert out.exists()
    res = runner.invoke(app, ["schema", "-o", str(tmp_path / "s")])
    assert res.exit_code == 0 and (tmp_path / "s" / "catalog.schema.json").exists()
