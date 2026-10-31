"""Sinh Tuyên bố áp dụng (SoA) dạng Excel và báo cáo khoảng cách dạng Markdown."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__
from .metrics import RULE, state_of, tally
from .models import STATUS_LABELS, Assessment, Catalog
from .repository import Repository

DISCLAIMER = (
    "Tài liệu do công cụ vn-compliance-crosswalk sinh ra từ dữ liệu tổ chức tự khai báo. "
    "Đối chiếu văn bản pháp luật là diễn giải chuyên môn, không phải ý kiến pháp lý."
)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
YES, NO = "Có", "Không"


def _related_text(repo: Repository, catalog: Catalog, control_id: str, targets: set[str] | None):
    items = repo.related(catalog.id, control_id, targets)
    return "; ".join(f"{r.catalog}:{r.control or '-'} ({r.relationship}, {r.confidence})" for r in items)


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_soa(
    repo: Repository,
    assessment: Assessment,
    out: str | Path,
    related_catalogs: set[str] | None = None,
    level: int | None = None,
) -> Path:
    level = level or assessment.level
    catalog = repo.catalog(assessment.catalog).for_level(level)
    wb = Workbook()

    # ---------------------------------------------------------------- Thông tin
    info = wb.active
    info.title = "Thông tin"
    rows = [
        ("Tổ chức", assessment.organization),
        ("Phạm vi", assessment.scope),
        ("Danh mục", f"{catalog.title} ({catalog.id}, phiên bản {catalog.version})"),
        ("Cấp độ HTTT", f"Cấp độ {level}" if level else "Không lọc theo cấp độ"),
        ("Ngày chốt số liệu", assessment.as_of.isoformat()),
        ("Ngày tạo tệp", date.today().isoformat()),
        ("Công cụ", f"vn-compliance-crosswalk {__version__}"),
        ("Quy tắc tính tỷ lệ", RULE),
        ("Lưu ý", DISCLAIMER),
    ]
    if catalog.notice:
        rows.append(("Lưu ý danh mục", catalog.notice))
    for r, (k, v) in enumerate(rows, 1):
        info.cell(row=r, column=1, value=k).font = Font(bold=True)
        info.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 100

    # ---------------------------------------------------------------- SoA
    ws = wb.create_sheet("SoA")
    headers = [
        "Mã",
        "Nhóm",
        "Biện pháp",
        "Áp dụng",
        "Lý do áp dụng / không áp dụng",
        "Trạng thái",
        "Bằng chứng",
        "Người phụ trách",
        "Hạn hoàn thành",
        "Yêu cầu liên quan",
        "Ghi chú",
    ]
    widths = [8, 24, 44, 10, 44, 22, 40, 18, 14, 40, 30]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for group, control in catalog.controls():
        s = state_of(assessment, control.id)
        ws.append(
            [
                control.id,
                f"{group.id} - {group.title}",
                control.title,
                YES if s.applicable else NO,
                s.justification or "",
                STATUS_LABELS[s.status] if s.applicable else "",
                "\n".join(s.evidence),
                s.owner or "",
                s.due.isoformat() if s.due else "",
                _related_text(repo, catalog, control.id, related_catalogs),
                s.notes or "",
            ]
        )
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"

    dv_app = DataValidation(type="list", formula1=f'"{YES},{NO}"', allow_blank=False)
    dv_status = DataValidation(
        type="list", formula1='"' + ",".join(STATUS_LABELS.values()) + '"', allow_blank=True
    )
    ws.add_data_validation(dv_app)
    ws.add_data_validation(dv_status)
    dv_app.add(f"D2:D{last}")
    dv_status.add(f"F2:F{last}")

    # ---------------------------------------------------------------- Tổng hợp (công thức sống)
    sm = wb.create_sheet("Tổng hợp")
    sm_headers = [
        "Nhóm",
        "Tổng",
        "Không áp dụng",
        "Áp dụng",
        "Đã triển khai",
        "Một phần",
        "Đã lập kế hoạch",
        "Chưa triển khai",
        "Chưa đánh giá",
        "Tỷ lệ tuân thủ",
    ]
    sm.append(sm_headers)
    _style_header(sm, 1, len(sm_headers))
    rng = lambda col: f"SoA!${col}$2:${col}${last}"  # noqa: E731
    group_names = [f"{g.id} - {g.title}" for g in catalog.groups]
    for r, name in enumerate(group_names, 2):
        crit = f"{rng('B')},$A{r}"
        status_count = lambda label, crit=crit: (  # noqa: E731
            f'=COUNTIFS({crit},{rng("D")},"{YES}",{rng("F")},"{label}")'
        )
        sm.append(
            [
                name,
                f"=COUNTIF({crit})",
                f'=COUNTIFS({crit},{rng("D")},"{NO}")',
                f"=B{r}-C{r}",
                status_count(STATUS_LABELS["implemented"]),
                status_count(STATUS_LABELS["partial"]),
                status_count(STATUS_LABELS["planned"]),
                status_count(STATUS_LABELS["not_implemented"]),
                status_count(STATUS_LABELS[None]),
                f'=IF(D{r}=0,"",(E{r}+0.5*F{r})/D{r})',
            ]
        )
    total_row = len(group_names) + 2
    sm.append(
        ["Toàn bộ"]
        + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{total_row - 1})" for c in range(2, 10)]
        + [f'=IF(D{total_row}=0,"",(E{total_row}+0.5*F{total_row})/D{total_row})']
    )
    for cell in sm[total_row]:
        cell.font = Font(bold=True)
    for r in range(2, total_row + 1):
        sm.cell(row=r, column=10).number_format = "0.0%"
    sm.column_dimensions["A"].width = 30
    for c in range(2, 11):
        sm.column_dimensions[get_column_letter(c)].width = 14
    sm.cell(row=total_row + 2, column=1, value="Quy tắc tính: " + RULE)

    out = Path(out)
    wb.save(out)
    return out


def gap_markdown(
    repo: Repository,
    assessment: Assessment,
    related_catalogs: set[str] | None = None,
    level: int | None = None,
) -> str:
    level = level or assessment.level
    catalog = repo.catalog(assessment.catalog).for_level(level)
    stats = tally(catalog, assessment)
    titles = {g.id: g.title for g in catalog.groups}
    lines = [
        f"# Báo cáo khoảng cách - {catalog.title}",
        "",
        f"- **Tổ chức:** {assessment.organization}",
        f"- **Phạm vi:** {assessment.scope}",
        f"- **Cấp độ HTTT:** {f'cấp độ {level}' if level else 'không lọc theo cấp độ'}",
        f"- **Ngày chốt số liệu:** {assessment.as_of.isoformat()}",
        f"- **Danh mục:** {catalog.id} (phiên bản {catalog.version}, trạng thái {catalog.status})",
        f"- **Quy tắc tính:** {RULE}",
        "",
        "## I. Tổng hợp theo nhóm",
        "",
        "| Nhóm | Tổng | Không áp dụng | Áp dụng (mẫu số) | Đã triển khai | Một phần "
        "| Kế hoạch | Chưa triển khai | Chưa đánh giá | Tỷ lệ |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for key, t in stats.items():
        if key == "*":
            continue
        lines.append(_tally_row(f"{key} - {titles[key]}", t))
    lines.append(_tally_row("**Toàn bộ**", stats["*"]))

    lines += ["", "## II. Danh sách khoảng cách", ""]
    gaps = []
    for _, control in catalog.controls():
        s = state_of(assessment, control.id)
        if s.applicable and s.status != "implemented":
            gaps.append((control, s))
    if gaps:
        lines += [
            "| Mã | Biện pháp | Trạng thái | Người phụ trách | Hạn | Yêu cầu liên quan |",
            "|---|---|---|---|---|---|",
        ]
        for control, s in gaps:
            rel = _related_text(repo, catalog, control.id, related_catalogs) or "-"
            lines.append(
                f"| {control.id} | {control.title} | {STATUS_LABELS[s.status]} "
                f"| {s.owner or '-'} | {s.due.isoformat() if s.due else '-'} | {rel} |"
            )
    else:
        lines.append("Không có khoảng cách: mọi control áp dụng đều đã triển khai.")

    excluded = [(c, state_of(assessment, c.id)) for _, c in catalog.controls()]
    excluded = [(c, s) for c, s in excluded if not s.applicable]
    lines += ["", "## III. Control không áp dụng và lý do", ""]
    if excluded:
        lines += [f"- **{c.id} {c.title}:** {s.justification}" for c, s in excluded]
    else:
        lines.append("Không có.")
    lines += ["", "---", f"*{DISCLAIMER}*", ""]
    return "\n".join(lines)


def _tally_row(label: str, t) -> str:
    c = t.counts
    score = "-" if t.score is None else f"{t.score:.1%}".replace(".", ",")
    return (
        f"| {label} | {t.total} | {t.not_applicable} | {t.applicable} | {c['implemented']} "
        f"| {c['partial']} | {c['planned']} | {c['not_implemented']} | {c['unassessed']} | {score} |"
    )


def assessment_template(catalog: Catalog, organization: str, scope: str, level: int | None = None) -> str:
    """Sinh tệp YAML hồ sơ đánh giá trống, đủ mọi control của danh mục (lọc theo cấp độ nếu có)."""
    catalog = catalog.for_level(level)
    lines = [
        "# Hồ sơ đánh giá - điền status: implemented | partial | planned | not_implemented",
        "# Control không áp dụng: applicable: false và bắt buộc ghi justification.",
        "# KHÔNG commit hồ sơ thật của tổ chức lên repo công khai.",
        "kind: assessment",
        f"organization: {organization}",
        f"scope: {scope}",
        f"catalog: {catalog.id}",
        f"as_of: {date.today().isoformat()}",
        f"level: {level if level else 'null'}",
        "controls:",
    ]
    for _, c in catalog.controls():
        lines += [
            f'  "{c.id}":  # {c.title}',
            "    applicable: true",
            "    status: null",
            "    owner: null",
            "    evidence: []",
        ]
    return "\n".join(lines) + "\n"


def coverage_markdown(cov) -> str:
    """Báo cáo độ phủ của một bảng đối chiếu."""
    m = cov.mapping
    pct = lambda a, b: f"{a / b:.0%}" if b else "-"  # noqa: E731
    lines = [
        f"# Độ phủ bảng đối chiếu {m.id}",
        "",
        f"- Phiên bản {m.version}, trạng thái {m.status}, {len(m.entries)} dòng đối chiếu.",
        f"- **{m.target_catalog}:** {cov.target_covered}/{cov.target_total} yêu cầu có ít nhất một "
        f"biện pháp đối chiếu ({pct(cov.target_covered, cov.target_total)}).",
        f"- **{m.source_catalog}:** {cov.source_covered}/{cov.source_total} biện pháp đã được xem xét "
        f"({pct(cov.source_covered, cov.source_total)}).",
        "- Quy tắc đếm: một mục được tính là đã phủ khi có ít nhất một dòng đối chiếu bất kỳ "
        "(kể cả no_relationship ở phía nguồn).",
        "",
        f"## I. Yêu cầu của {m.target_catalog} chưa có biện pháp đối chiếu",
        "",
        "Đây thường là nghĩa vụ hành chính, pháp lý đặc thù, cần biện pháp riêng ngoài bộ biện pháp nguồn.",
        "",
    ]
    lines += [f"- **{c.id}** ({g}) {c.title}" for g, c in cov.target_uncovered] or ["Không có."]
    lines += ["", f"## II. Biện pháp của {m.source_catalog} chưa được xem xét", ""]
    lines += [f"- **{c.id}** ({g}) {c.title}" for g, c in cov.source_uncovered] or ["Không có."]
    return "\n".join(lines) + "\n"
