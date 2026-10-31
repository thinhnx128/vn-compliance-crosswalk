"""Xuất gói Excel cho người dùng không dùng dòng lệnh.

Gói gồm: hướng dẫn; bảng tự đánh giá có ô chọn cấp độ (công thức sống); danh mục yêu cầu đích;
bảng đối chiếu; danh mục nguồn kèm yêu cầu liên quan; độ phủ; tiêu chí xác định cấp độ.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__
from .levels import criteria_data
from .metrics import RULE
from .models import LEVEL_TAG, STATUS_LABELS, applies_to_level
from .reports import DISCLAIMER
from .repository import Repository

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FONT = Font(bold=True, size=14)
ROLE_LABELS = {
    "chu-quan": "Chủ quản",
    "don-vi-chuyen-trach": "Đơn vị chuyên trách",
    "don-vi-van-hanh": "Đơn vị vận hành",
}
VERIFY_TAG = "can-doi-chieu"
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row: int, headers: list[str], widths: list[int]) -> None:
    for i, (h, w) in enumerate(zip(headers, widths, strict=True), 1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w


def _roles(tags: list[str]) -> str:
    return ", ".join(ROLE_LABELS[t] for t in tags if t in ROLE_LABELS)


def _related(repo: Repository, catalog_id: str, control_id: str, other: str) -> str:
    return "; ".join(
        f"{r.control} ({r.relationship}, {r.confidence})"
        for r in repo.related(catalog_id, control_id, {other})
        if r.control
    )


def export_xlsx(repo: Repository, mapping_id: str, out: str | Path) -> Path:
    mapping = repo.mappings[mapping_id]
    src = repo.catalog(mapping.source_catalog)
    dst = repo.catalog(mapping.target_catalog)
    cov = repo.coverage(mapping_id)
    wb = Workbook()

    # ------------------------------------------------------------------ Hướng dẫn
    ws = wb.active
    ws.title = "Hướng dẫn"
    ws["A1"] = f"Gói tự đánh giá {dst.title}"
    ws["A1"].font = TITLE_FONT
    rows = [
        ("Phiên bản gói", f"vn-compliance-crosswalk {__version__}, xuất ngày {date.today().isoformat()}"),
        (
            "Văn bản",
            f"{dst.document_number or dst.title}, ban hành {dst.issued_on or '-'}, hiệu lực {dst.effective_date or '-'}",
        ),
        ("Nguồn đối chiếu", dst.source_url or "-"),
        (
            "Trạng thái dữ liệu",
            f"Danh mục: {dst.status}. Bảng đối chiếu: {mapping.status}, {len(mapping.entries)} dòng.",
        ),
        (
            "Cách dùng",
            "1. Sang trang 'Tự đánh giá', chọn cấp độ HTTT ở ô màu vàng (dùng trang 'Tiêu chí cấp độ' "
            "hoặc công cụ web để xác định). 2. Cột 'Áp dụng' tự cập nhật theo cấp độ. 3. Chọn trạng thái, "
            "ghi bằng chứng, người phụ trách, hạn cho từng yêu cầu áp dụng. 4. Bảng tổng hợp ở đầu trang tự tính.",
        ),
        ("Quy tắc tính", RULE),
        (
            "Thẻ 'Cần đối chiếu'",
            "Yêu cầu tóm tắt chưa đối chiếu được toàn văn; phải kiểm tra văn bản gốc trước khi dùng.",
        ),
        ("Công cụ web", "https://thinhnx128.github.io/vn-compliance-crosswalk/cap-do/"),
        ("Mã nguồn", "https://github.com/thinhnx128/vn-compliance-crosswalk"),
        ("Giấy phép", "Dữ liệu CC BY 4.0, mã nguồn MIT. Không chứa nội dung tiêu chuẩn ISO."),
        ("Lưu ý", DISCLAIMER),
    ]
    if dst.notice:
        rows.append(("Lưu ý danh mục", dst.notice))
    if mapping.notice:
        rows.append(("Lưu ý đối chiếu", mapping.notice))
    for r, (k, v) in enumerate(rows, 3):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v).alignment = WRAP_TOP
        ws.cell(row=r, column=1).alignment = WRAP_TOP
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110

    # ------------------------------------------------------------------ Tự đánh giá
    sa = wb.create_sheet("Tự đánh giá")
    sa["A1"] = f"Tự đánh giá tuân thủ {dst.document_number or dst.title}"
    sa["A1"].font = TITLE_FONT
    inputs = [
        ("Cấp độ HTTT (1–5)", 2),
        ("Tổ chức", ""),
        ("Hệ thống / phạm vi", ""),
        ("Ngày chốt số liệu", ""),
    ]
    for i, (k, v) in enumerate(inputs, 2):
        sa.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = sa.cell(row=i, column=3, value=v)
        c.fill = INPUT_FILL
    level_cell = "$C$2"
    dv_level = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=False)
    dv_level.error, dv_level.errorTitle = "Nhập số từ 1 đến 5", "Cấp độ không hợp lệ"
    sa.add_data_validation(dv_level)
    dv_level.add("C2")

    headers = [
        "Mã",
        "Nhóm",
        "Yêu cầu",
        "Chủ thể",
        "Cần đối chiếu",
        "Áp dụng",
        "Trạng thái",
        "Bằng chứng",
        "Người phụ trách",
        "Hạn",
        "Ghi chú",
        f"Biện pháp {src.id} liên quan",
        "CĐ1",
        "CĐ2",
        "CĐ3",
        "CĐ4",
        "CĐ5",
    ]
    widths = [11, 26, 56, 20, 9, 9, 20, 34, 18, 12, 26, 30, 5, 5, 5, 5, 5]
    head_row = 16
    _header(sa, head_row, headers, widths)
    first = head_row + 1
    controls = list(dst.controls())
    last = first + len(controls) - 1
    col = {h: get_column_letter(i) for i, h in enumerate(headers, 1)}
    lv_first, lv_last = col["CĐ1"], col["CĐ5"]
    for r, (group, c) in enumerate(controls, first):
        flags = [1 if applies_to_level(c, n) else 0 for n in range(1, 6)]
        sa.append(
            [
                c.id,
                f"{group.id} - {group.title}",
                c.title + (f"\n{c.summary.strip()}" if c.summary else ""),
                _roles(c.tags),
                "Có" if VERIFY_TAG in c.tags else "",
                f'=IF(INDEX(${lv_first}{r}:${lv_last}{r},1,{level_cell})=1,"Có","Không")',
                "",
                "",
                "",
                "",
                "",
                _related(repo, dst.id, c.id, src.id),
                *flags,
            ]
        )
        for cell in sa[r]:
            cell.alignment = WRAP_TOP
    for h in ("CĐ1", "CĐ2", "CĐ3", "CĐ4", "CĐ5"):
        sa.column_dimensions[col[h]].hidden = True
    statuses = [v for k, v in STATUS_LABELS.items() if k is not None]
    dv_status = DataValidation(type="list", formula1='"' + ",".join(statuses) + '"', allow_blank=True)
    sa.add_data_validation(dv_status)
    dv_status.add(f"{col['Trạng thái']}{first}:{col['Trạng thái']}{last}")
    sa.conditional_formatting.add(
        f"A{first}:{col['Biện pháp ' + src.id + ' liên quan']}{last}",
        FormulaRule(formula=[f'${col["Áp dụng"]}{first}="Không"'], font=Font(color="A6A6A6")),
    )
    sa.freeze_panes = f"D{first}"
    sa.auto_filter.ref = f"A{head_row}:{col['Biện pháp ' + src.id + ' liên quan']}{last}"

    ap = f"${col['Áp dụng']}${first}:${col['Áp dụng']}${last}"
    st = f"${col['Trạng thái']}${first}:${col['Trạng thái']}${last}"
    summary = [
        ("Số yêu cầu áp dụng (mẫu số)", f'=COUNTIF({ap},"Có")'),
        (STATUS_LABELS["implemented"], f'=COUNTIFS({ap},"Có",{st},"{STATUS_LABELS["implemented"]}")'),
        (STATUS_LABELS["partial"], f'=COUNTIFS({ap},"Có",{st},"{STATUS_LABELS["partial"]}")'),
        (STATUS_LABELS["planned"], f'=COUNTIFS({ap},"Có",{st},"{STATUS_LABELS["planned"]}")'),
        (STATUS_LABELS["not_implemented"], f'=COUNTIFS({ap},"Có",{st},"{STATUS_LABELS["not_implemented"]}")'),
        (STATUS_LABELS[None], "=C7-SUM(C8:C11)"),
        ("Tỷ lệ tuân thủ", '=IF(C7=0,"",(C8+0.5*C9)/C7)'),
    ]
    for i, (k, f) in enumerate(summary, 7):
        sa.cell(row=i, column=1, value=k).font = Font(bold=True)
        sa.cell(row=i, column=3, value=f)
    sa["C13"].number_format = "0.0%"
    sa.cell(row=14, column=1, value="Quy tắc tính: " + RULE).font = Font(italic=True, color="595959")

    # ------------------------------------------------------------------ Danh mục đích
    dc = wb.create_sheet(f"Danh mục {dst.id}"[:31])
    h2 = [
        "Mã",
        "Nhóm",
        "Yêu cầu",
        "Tóm tắt",
        "Chủ thể",
        "Cấp độ 1",
        "Cấp độ 2",
        "Cấp độ 3",
        "Cấp độ 4",
        "Cấp độ 5",
        "Cần đối chiếu",
    ]
    _header(dc, 1, h2, [11, 28, 56, 56, 22, 9, 9, 9, 9, 9, 10])
    for group, c in dst.controls():
        dc.append(
            [
                c.id,
                f"{group.id} - {group.title}",
                c.title,
                (c.summary or "").strip(),
                _roles(c.tags),
                *["x" if applies_to_level(c, n) else "" for n in range(1, 6)],
                "Có" if VERIFY_TAG in c.tags else "",
            ]
        )
    for row in dc.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
    dc.freeze_panes = "C2"
    dc.auto_filter.ref = dc.dimensions

    # ------------------------------------------------------------------ Đối chiếu
    mp = wb.create_sheet("Đối chiếu")
    s_idx, d_idx = src.index(), dst.index()
    h3 = [
        f"Mã {src.id}",
        "Biện pháp",
        f"Mã {dst.id}",
        "Yêu cầu",
        "Quan hệ",
        "Mức tin cậy",
        "Lý do",
        "Người rà soát",
        "Ngày rà soát",
    ]
    _header(mp, 1, h3, [10, 36, 12, 46, 14, 11, 70, 14, 12])
    for e in mapping.entries:
        mp.append(
            [
                e.source,
                s_idx[e.source].title,
                e.target or "-",
                d_idx[e.target].title if e.target else "-",
                e.relationship,
                e.confidence,
                e.rationale,
                e.reviewed_by or "",
                e.reviewed_on.isoformat() if e.reviewed_on else "",
            ]
        )
    for row in mp.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
    mp.freeze_panes = "B2"
    mp.auto_filter.ref = mp.dimensions

    # ------------------------------------------------------------------ Danh mục nguồn
    sc = wb.create_sheet(f"Danh mục {src.id}"[:31])
    _header(
        sc,
        1,
        ["Mã", "Nhóm", "Biện pháp", "Tên tiếng Anh", f"Yêu cầu {dst.id} liên quan"],
        [9, 26, 46, 44, 50],
    )
    for group, c in src.controls():
        sc.append(
            [
                c.id,
                f"{group.id} - {group.title}",
                c.title,
                c.title_en or "",
                _related(repo, src.id, c.id, dst.id),
            ]
        )
    for row in sc.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
    sc.freeze_panes = "C2"
    sc.auto_filter.ref = sc.dimensions

    # ------------------------------------------------------------------ Độ phủ
    cv = wb.create_sheet("Độ phủ")
    cv["A1"] = f"Độ phủ bảng đối chiếu {mapping.id}"
    cv["A1"].font = TITLE_FONT
    cv["A3"], cv["B3"] = (
        f"Yêu cầu {dst.id} có biện pháp đối chiếu",
        f"{cov.target_covered}/{cov.target_total}",
    )
    cv["A4"], cv["B4"] = f"Biện pháp {src.id} đã xem xét", f"{cov.source_covered}/{cov.source_total}"
    cv["A5"] = "Quy tắc: một mục được tính là đã phủ khi có ít nhất một dòng đối chiếu."
    cv["A7"] = (
        f"Yêu cầu {dst.id} chưa có biện pháp tương ứng (thường là nghĩa vụ hành chính, pháp lý đặc thù)"
    )
    cv["A7"].font = Font(bold=True)
    _header(cv, 8, ["Mã", "Nhóm", "Yêu cầu"], [12, 12, 90])
    for g, c in cov.target_uncovered:
        cv.append([c.id, g, c.title])
    for k in ("A3", "A4", "A5"):
        cv[k].alignment = WRAP_TOP

    # ------------------------------------------------------------------ Tiêu chí cấp độ
    if dst.id.startswith("nd331"):
        data = criteria_data()
        tc = wb.create_sheet("Tiêu chí cấp độ")
        tc["A1"] = data["title"]
        tc["A1"].font = TITLE_FONT
        _header(
            tc, 3, ["Cấp độ", "Căn cứ", "Tiêu chí (đáp ứng một trong các tiêu chí của cấp độ)"], [9, 22, 100]
        )
        for c in sorted(data["criteria"], key=lambda x: (x["level"], x["id"])):
            tc.append([c["level"], c["article"], c["label"]])
        r = tc.max_row + 2
        tc.cell(row=r, column=1, value="Thẩm quyền").font = Font(bold=True)
        for lv in sorted(data["approval"]):
            r += 1
            tc.cell(row=r, column=1, value=int(lv))
            tc.cell(row=r, column=3, value=data["approval"][lv])
        r += 2
        tc.cell(row=r, column=1, value="Lưu ý").font = Font(bold=True)
        for t in data["reminders"]:
            r += 1
            tc.cell(row=r, column=3, value=t)
        r += 2
        tc.cell(row=r, column=3, value=data["notice"]).font = Font(italic=True, color="595959")
        for row in tc.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = WRAP_TOP

    wb.active = 1  # mở thẳng trang Tự đánh giá
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


__all__ = ["export_xlsx", "LEVEL_TAG"]
